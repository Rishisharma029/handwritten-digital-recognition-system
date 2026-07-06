"""DOCX (Word) export functionality."""
import os
from datetime import datetime
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import logging

logger = logging.getLogger(__name__)


class DocxExporter:
    """Export OCR results to DOCX (Word) format."""
    
    def __init__(self, output_dir: str = "outputs/docx"):
        """
        Initialize DOCX exporter.
        
        Args:
            output_dir (str): Output directory for DOCX files
        """
        self.output_dir = output_dir
        os.makedirs(self.output_dir, exist_ok=True)
    
    def export(
        self,
        filename: str,
        extracted_text: str,
        confidence: float,
        engine: str = "TrOCR",
        language: str = "English",
        metadata: dict = None
    ) -> str:
        """
        Export OCR result to DOCX file.
        
        Args:
            filename (str): Original filename
            extracted_text (str): Recognized text
            confidence (float): Confidence score
            engine (str): OCR engine used
            language (str): Language detected
            metadata (dict): Additional metadata
            
        Returns:
            str: Path to exported file
        """
        docx_filename = os.path.splitext(filename)[0] + ".docx"
        output_path = os.path.join(self.output_dir, docx_filename)
        
        try:
            doc = Document()
            
            # Title
            title = doc.add_heading("OCR Result Document", 0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            # Metadata section
            doc.add_heading("Document Information", level=2)
            
            table = doc.add_table(rows=5, cols=2)
            table.style = 'Light Grid Accent 1'
            
            cells = table.rows[0].cells
            cells[0].text = "Original File"
            cells[1].text = filename
            
            cells = table.rows[1].cells
            cells[0].text = "OCR Engine"
            cells[1].text = engine
            
            cells = table.rows[2].cells
            cells[0].text = "Language"
            cells[1].text = language
            
            cells = table.rows[3].cells
            cells[0].text = "Confidence"
            cells[1].text = f"{confidence:.2%}"
            
            cells = table.rows[4].cells
            cells[0].text = "Processed"
            cells[1].text = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
            
            # Text content
            doc.add_heading("Recognized Text", level=2)
            doc.add_paragraph(extracted_text)
            
            # Save document
            doc.save(output_path)
            logger.info(f"DOCX exported to {output_path}")
            return output_path
            
        except Exception as e:
            logger.error(f"Error exporting DOCX: {e}")
            raise
